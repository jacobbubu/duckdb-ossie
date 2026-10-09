"""Read-only source extraction. Run with the bundled workspace Python runtime."""

import hashlib
import json
from pathlib import Path

import openpyxl
from docx import Document


ROOT = Path(__file__).resolve().parent
MANIFEST = ROOT / "sources" / "manifest.json"


def main():
    manifest = json.loads(MANIFEST.read_text())
    extracted = {}
    for entry in manifest["files"]:
        source = Path(entry["source_path"])
        if source.suffix not in {".xlsx", ".docx"}:
            continue
        actual = hashlib.sha256(source.read_bytes()).hexdigest()
        if actual != entry["sha256"]:
            raise ValueError(f"Source changed: {source}")
        record = {"source": str(source), "source_sha256": actual}
        if source.suffix == ".xlsx":
            workbook = openpyxl.load_workbook(source, read_only=True, data_only=False)
            record["sheets"] = {
                sheet.title: [
                    {cell.coordinate: cell.value for cell in row if cell.value is not None}
                    for row in sheet.iter_rows()
                ]
                for sheet in workbook
            }
            workbook.close()
        else:
            document = Document(source)
            record["paragraphs"] = [
                {"index": i, "text": paragraph.text}
                for i, paragraph in enumerate(document.paragraphs)
                if paragraph.text
            ]
            record["tables"] = [
                [[cell.text for cell in row.cells] for row in table.rows]
                for table in document.tables
            ]
        extracted[source.name] = record
    out = ROOT / "sources" / "business-source-extraction.json"
    if out.exists() and json.loads(out.read_text()) != extracted:
        raise ValueError("Refusing to overwrite a different raw extraction")
    out.write_text(json.dumps(extracted, ensure_ascii=False, indent=2) + "\n")
    print(f"Extracted {len(extracted)} hash-verified business sources")


if __name__ == "__main__":
    main()
